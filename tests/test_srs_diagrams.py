"""
test_srs_diagrams.py — requirements read out of figures in the SRS
==================================================================
Figures (UML, flowcharts, architecture diagrams) are emitted by the .docx and
.pdf readers and interpreted by src/srs/diagrams.py. Nothing here touches the
network: the vision model is tests.fakes.FakeVisionClient and OCR is a
patched pytesseract.
"""

import io
import json
from pathlib import Path

import pytest
from docx import Document
from PIL import Image, ImageDraw

from src.config import ConfigError, config_from_dict
from src.srs import ExtractionOptions, extract_requirements_with_report
from src.srs import diagrams
from src.srs.blocks import Figure, Paragraph
from src.srs.diagrams import (
    AltTextEngine,
    TesseractEngine,
    VisionEngine,
    attach_captions,
    figure_label,
    prepare_image,
    statements_from_text,
)
from src.srs.options import DEFAULT_ID_PATTERN
from src.srs.readers.docx_reader import read_docx
from src.srs.readers.pdf_reader import read_pdf
from src.srs.strategies import TagMatcher
from tests.fakes import FakeVisionClient

ROOT = Path(__file__).parent.parent
FIXTURES = ROOT / "tests" / "fixtures" / "srs"
TAGS = TagMatcher(DEFAULT_ID_PATTERN)

LOCAL = ExtractionOptions(diagram_mode="local")
VISION = ExtractionOptions(diagram_mode="vision")


def _png(w=400, h=240, color="white", text="Customer -> Place Order") -> bytes:
    img = Image.new("RGB", (w, h), color)
    ImageDraw.Draw(img).text((10, 10), text, fill="black")
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    return buf.getvalue()


def _docx(tmp_path, items, name="srs.docx") -> Path:
    """items: ("h1"/"h2"/"h3", text) | ("p", text) | ("caption", text) | ("table", rows)
    | ("img", alt_text, png_bytes or None)."""
    doc = Document()
    for item in items:
        kind = item[0]
        if kind.startswith("h"):
            doc.add_heading(item[1], level=int(kind[1]))
        elif kind == "p":
            doc.add_paragraph(item[1])
        elif kind == "caption":
            doc.add_paragraph(item[1], style="Caption")
        elif kind == "table":
            rows = item[1]
            t = doc.add_table(rows=len(rows), cols=len(rows[0]))
            for r, row in enumerate(rows):
                for c, val in enumerate(row):
                    t.cell(r, c).text = val
        elif kind == "img":
            shape = doc.add_paragraph().add_run().add_picture(io.BytesIO(item[2] or _png()))
            shape._inline.docPr.set("descr", item[1])
    path = tmp_path / name
    doc.save(path)
    return path


def _use_vision(monkeypatch, client):
    monkeypatch.setattr(diagrams, "vision_client_factory", lambda: client)


REQ_TABLE = ("table", [
    ("ID", "Title", "Description"),
    ("REQ-01", "User login", "The system must authenticate users with a password and return a JWT."),
    ("REQ-04", "Refund integrity", "Refunds must include a client-provided transaction_id for idempotency."),
])


# ---------------------------------------------------------------------------
# Readers
# ---------------------------------------------------------------------------

class TestReaders:
    def test_docx_emits_figure_with_alt_text_and_bytes(self, tmp_path):
        path = _docx(tmp_path, [("h1", "Overview"), ("img", "Context diagram of the store", None),
                                ("caption", "Figure 1: Context")])
        figs = [b for b in read_docx(path) if isinstance(b, Figure)]
        assert len(figs) == 1
        assert figs[0].content_type == "image/png" and figs[0].image.startswith(b"\x89PNG")
        assert figs[0].alt_text == "Context diagram of the store"
        assert str(figs[0].locator) == "figure 1"

    def test_word_auto_alt_text_suffix_is_removed(self, tmp_path):
        path = _docx(tmp_path, [("img", "A diagram of a checkout flow. Description automatically generated", None)])
        fig = next(b for b in read_docx(path) if isinstance(b, Figure))
        assert fig.alt_text == "A diagram of a checkout flow."

    def test_pdf_renders_images_as_figures(self):
        figs = [b for b in read_pdf(FIXTURES / "ecommerce_srs_ieee.pdf") if isinstance(b, Figure)]
        assert len(figs) == 1
        assert figs[0].image.startswith(b"\x89PNG") and figs[0].locator.page is not None


# ---------------------------------------------------------------------------
# Building blocks
# ---------------------------------------------------------------------------

class TestHelpers:
    def test_caption_after_or_before(self):
        from src.srs.blocks import BlockKind, Locator

        fig = Figure(locator=Locator(BlockKind.FIGURE, 3), image=b"")
        cap = Paragraph(text="Figure 2: Refund sequence", locator=Locator(BlockKind.PARAGRAPH, 1), style="Caption")
        after = attach_captions([fig, cap])[0]
        before = attach_captions([cap, fig])[1]
        assert after.caption == before.caption == "Figure 2: Refund sequence"
        assert figure_label(after) == "FIG-2"
        assert figure_label(fig) == "FIG-3"                     # no caption: document ordinal

    def test_prepare_image(self):
        ok, reason = prepare_image(_png())
        assert ok is not None and reason == "" and ok.png.startswith(b"\x89PNG")
        assert prepare_image(_png(16, 16))[1] == "image_too_small"
        assert prepare_image(b"not an image")[1] == "image_unreadable"
        big, _ = prepare_image(_png(5000, 1000))
        assert max(big.width, big.height) == diagrams.MAX_SIDE_PX

    def test_transparent_png_is_flattened_on_white(self):
        img = Image.new("RGBA", (200, 200), (0, 0, 0, 0))
        buf = io.BytesIO()
        img.save(buf, format="PNG")
        prepared, _ = prepare_image(buf.getvalue())
        with Image.open(io.BytesIO(prepared.png)) as out:
            assert out.mode == "RGB" and out.getpixel((5, 5)) == (255, 255, 255)

    def test_statements_from_text(self):
        statements, labels = statements_from_text(
            "Customer\nPlace Order\n\nREQ-12: The system shall email a receipt. "
            "The system must log every refund attempt.", TAGS)
        assert [(s.req_id, s.text) for s in statements] == [
            ("REQ-12", "The system shall email a receipt."),
            ("", "The system must log every refund attempt."),
        ]
        assert labels == ["Customer Place Order"]


# ---------------------------------------------------------------------------
# Local mode (alt text, OCR)
# ---------------------------------------------------------------------------

class TestLocalMode:
    def test_tagged_alt_text_becomes_a_requirement(self, tmp_path):
        path = _docx(tmp_path, [
            ("h1", "3 Specific Requirements"), REQ_TABLE,
            ("h2", "3.2 Security workflows"),
            ("img", "REQ-09: The system shall require 2FA verification for orders above $500.", None),
            ("caption", "Figure 4: High-value checkout activity"),
        ])
        reqs, report = extract_requirements_with_report(path, LOCAL)
        by_id = {r["id"]: r for r in reqs}
        assert list(by_id) == ["REQ-01", "REQ-04", "REQ-09"]
        req = by_id["REQ-09"]
        assert req["strategy"] == "diagram" and req["id_source"] == "document"
        assert req["text"] == "The system shall require 2FA verification for orders above $500."
        assert req["section"] == "3 Specific Requirements > 3.2 Security workflows"
        assert req["source"]["locator"] == "figure 1"
        assert req["figures"][0]["figure"] == "FIG-4" and req["figures"][0]["engine"] == "alt_text"
        assert report.by_strategy == {"table": 2, "diagram": 1}
        assert report.figures[0]["requirements"] == ["REQ-09"]

    def test_untagged_statements_get_generated_ids(self, tmp_path):
        path = _docx(tmp_path, [
            ("h1", "Requirements"), REQ_TABLE,
            ("h2", "Order lifecycle"),
            ("img", "The system shall cancel unpaid orders after 30 minutes. "
                    "The system must notify the customer when an order ships.", None),
            ("caption", "Figure 7: Order state machine"),
        ])
        reqs, report = extract_requirements_with_report(path, LOCAL)
        gen = [r for r in reqs if r["strategy"] == "diagram"]
        assert [r["id"] for r in gen] == ["FIG-7.1", "FIG-7.2"]
        assert all(r["id_source"] == "generated" for r in gen)
        assert any("generated IDs (FIG-" in w for w in report.warnings)

    def test_figure_restating_a_known_requirement_is_context(self, tmp_path):
        path = _docx(tmp_path, [
            ("h1", "Requirements"), REQ_TABLE,
            ("img", "REQ-04: The system shall reject refunds without a transaction_id.", None),
        ])
        reqs, report = extract_requirements_with_report(path, LOCAL)
        assert [r["id"] for r in reqs] == ["REQ-01", "REQ-04"]
        refund = reqs[1]
        assert refund["strategy"] == "table"                       # the text version wins
        assert "transaction_id for idempotency" in refund["text"]  # and is not changed by the figure
        assert refund["figures"][0]["figure"] == "FIG-1"
        assert report.figures[0]["illustrates"] == ["REQ-04"] and not report.duplicates

    def test_figure_under_a_tagged_heading_illustrates_it(self, tmp_path):
        path = _docx(tmp_path, [
            ("h1", "Requirements"),
            ("h2", "REQ-07 Order tracking"),
            ("p", "The system must return the fulfilment status and timeline of an order."),
            ("img", "The system shall show a timeline with dispatched and delivered events.", None),
        ])
        reqs, report = extract_requirements_with_report(path, LOCAL)
        assert [r["id"] for r in reqs] == ["REQ-07"]
        assert reqs[0]["figures"][0]["figure"] == "FIG-1"
        assert report.figures[0]["illustrates"] == ["REQ-07"]

    def test_caption_id_claims_untagged_statements(self, tmp_path):
        path = _docx(tmp_path, [
            ("h1", "Requirements"), REQ_TABLE,
            ("img", "The system shall issue a time-limited OTP. The system shall verify the OTP before payment.", None),
            ("caption", "Figure 2: REQ-09 two-factor authentication flow"),
        ])
        reqs, _ = extract_requirements_with_report(path, LOCAL)
        req = next(r for r in reqs if r["id"] == "REQ-09")
        assert req["text"] == ("The system shall issue a time-limited OTP. "
                               "The system shall verify the OTP before payment.")

    def test_labels_only_are_context_not_requirements(self, tmp_path):
        path = _docx(tmp_path, [("h1", "Requirements"), REQ_TABLE, ("img", "Customer, Store API, Bank", None)])
        reqs, report = extract_requirements_with_report(path, LOCAL)
        assert [r["id"] for r in reqs] == ["REQ-01", "REQ-04"]
        assert report.figures[0]["summary"] == "Customer, Store API, Bank"
        assert report.figures[0]["requirements"] == []

    def test_unreadable_figure_is_reported_not_fatal(self, tmp_path, monkeypatch):
        monkeypatch.setattr(diagrams, "tesseract_available", lambda: False)
        path = _docx(tmp_path, [("h1", "Requirements"), REQ_TABLE, ("img", "", None)])
        reqs, report = extract_requirements_with_report(path, LOCAL)
        assert len(reqs) == 2 and report.skipped["image"] == 1 and report.figures == []
        assert any("Tesseract" in w and "SRS_DIAGRAM_MODE=vision" in w for w in report.warnings)

    def test_appendix_figures_are_never_interpreted(self, tmp_path, monkeypatch):
        client = FakeVisionClient()
        _use_vision(monkeypatch, client)
        path = _docx(tmp_path, [
            ("h1", "Requirements"), REQ_TABLE,
            ("h1", "Appendix A: Screenshots"), ("img", "REQ-77: The system shall do appendix things.", None),
        ])
        reqs, report = extract_requirements_with_report(path, VISION)
        assert [r["id"] for r in reqs] == ["REQ-01", "REQ-04"]
        assert client.calls == [] and report.figures == []

    def test_off_mode_ignores_figures(self, tmp_path):
        path = _docx(tmp_path, [("h1", "Requirements"), REQ_TABLE,
                                ("img", "REQ-09: The system shall require 2FA for big orders.", None)])
        reqs, report = extract_requirements_with_report(path, ExtractionOptions(diagram_mode="off"))
        assert [r["id"] for r in reqs] == ["REQ-01", "REQ-04"]
        assert report.skipped["figure_ignored"] == 1

    def test_max_figures(self, tmp_path):
        items = [("h1", "Requirements"), REQ_TABLE] + [
            ("img", f"REQ-1{i}: The system shall support feature number {i}.", None) for i in range(3)]
        reqs, report = extract_requirements_with_report(_docx(tmp_path, items),
                                                        ExtractionOptions(diagram_mode="local", max_figures=2))
        assert [r["id"] for r in reqs][-2:] == ["REQ-10", "REQ-11"]
        assert report.skipped["figure_over_limit"] == 1
        assert any("max_figures" in w for w in report.warnings)

    def test_ocr_engine(self, monkeypatch):
        pytesseract = pytest.importorskip("pytesseract")
        monkeypatch.setattr(pytesseract, "image_to_string", lambda *a, **k:
                            "Customer\n\nCheckout\n\nREQ-21: The system shall lock the cart during payment.\n")
        image, _ = prepare_image(_png())
        from src.srs.blocks import BlockKind, Locator

        reading = TesseractEngine(TAGS).read(Figure(locator=Locator(BlockKind.FIGURE, 1), image=b""), image, "")
        assert [(s.req_id, s.text) for s in reading.statements] == [
            ("REQ-21", "The system shall lock the cart during payment.")]
        assert reading.summary == "Text in the figure: Customer; Checkout"

    def test_alt_text_engine_without_alt_text(self):
        from src.srs.blocks import BlockKind, Locator

        assert AltTextEngine(TAGS).read(Figure(locator=Locator(BlockKind.FIGURE, 1), image=b""), None, "") is None


# ---------------------------------------------------------------------------
# Vision mode
# ---------------------------------------------------------------------------

ACTIVITY = {
    "diagram_type": "activity",
    "summary": "High-value checkout requires a one-time password before the order is placed.",
    "requirements": [
        {"id": "REQ-09", "title": "High-value checkout 2FA",
         "statement": "The system shall require 2FA verification for orders above $500"},
        {"id": "REQ-09", "title": "",
         "statement": "The system shall issue a time-limited OTP via POST /api/v1/auth/2fa/challenge"},
        {"id": "made-up id", "title": "Audit",
         "statement": "The system shall log every failed OTP attempt"},
        {"id": None, "statement": "too short"},
    ],
}


class TestVisionMode:
    def test_vision_reading_end_to_end(self, tmp_path, monkeypatch):
        client = FakeVisionClient(ACTIVITY)
        _use_vision(monkeypatch, client)
        path = _docx(tmp_path, [("h1", "Requirements"), REQ_TABLE, ("img", "", None),
                                ("caption", "Figure 3: Checkout activity diagram")])
        reqs, report = extract_requirements_with_report(path, VISION)
        by_id = {r["id"]: r for r in reqs}
        assert list(by_id) == ["REQ-01", "REQ-04", "REQ-09", "FIG-3.1"]
        assert by_id["REQ-09"]["title"] == "High-value checkout 2FA"
        assert by_id["REQ-09"]["text"] == (
            "High-value checkout 2FA. The system shall require 2FA verification for orders above $500. "
            "The system shall issue a time-limited OTP via POST /api/v1/auth/2fa/challenge.")
        # an ID that does not match the document's ID pattern is never kept
        assert by_id["FIG-3.1"]["text"] == "Audit. The system shall log every failed OTP attempt"
        fig = report.figures[0]
        assert fig["engine"] == "vision:gpt-4o-mini" and fig["diagram_type"] == "activity"
        assert fig["requirements"] == ["REQ-09", "FIG-3.1"]

        (call,) = client.calls
        assert call["temperature"] == 0 and call["response_format"] == {"type": "json_object"}
        image_part = call["messages"][1]["content"][1]
        assert image_part["image_url"]["url"].startswith("data:image/png;base64,")
        assert "Figure 3: Checkout activity diagram" in call["messages"][1]["content"][0]["text"]

    def test_readings_are_cached_so_reparsing_is_deterministic(self, tmp_path, monkeypatch):
        client = FakeVisionClient(ACTIVITY)
        _use_vision(monkeypatch, client)
        monkeypatch.setenv("SRS_DIAGRAM_CACHE_DIR", str(tmp_path / "cache"))
        path = _docx(tmp_path, [("h1", "Requirements"), REQ_TABLE, ("img", "", _png(text="cache me"))])
        first, _ = extract_requirements_with_report(path, VISION)
        second, _ = extract_requirements_with_report(path, VISION)
        assert first == second and len(client.calls) == 1
        assert len(list((tmp_path / "cache").glob("*.json"))) == 1

    def test_failure_falls_back_to_alt_text(self, tmp_path, monkeypatch):
        _use_vision(monkeypatch, FakeVisionClient(error=TimeoutError("read timed out")))
        path = _docx(tmp_path, [("h1", "Requirements"), REQ_TABLE,
                                ("img", "REQ-09: The system shall require 2FA for orders above $500.", None)])
        reqs, report = extract_requirements_with_report(path, VISION)
        assert reqs[-1]["id"] == "REQ-09" and reqs[-1]["figures"][0]["engine"] == "alt_text"
        assert any("TimeoutError" in w for w in report.warnings)

    def test_no_api_key_warns_and_reads_locally(self, tmp_path, monkeypatch):
        _use_vision(monkeypatch, None)
        path = _docx(tmp_path, [("h1", "Requirements"), REQ_TABLE,
                                ("img", "REQ-09: The system shall require 2FA for orders above $500.", None)])
        reqs, report = extract_requirements_with_report(path, VISION)
        assert reqs[-1]["id"] == "REQ-09"
        assert any("no OPENAI_API_KEY" in w for w in report.warnings)

    def test_decorative_image_adds_nothing_and_parity_holds(self, monkeypatch):
        """The IEEE fixture's context diagram specifies no behaviour: vision mode
        must give exactly the same 10 requirements as the text-only parser."""
        client = FakeVisionClient({"diagram_type": "architecture", "summary": "Client app talks to Nexus API.",
                                   "requirements": []})
        _use_vision(monkeypatch, client)
        local, _ = extract_requirements_with_report(FIXTURES / "ecommerce_srs_ieee.docx", LOCAL)
        vision, report = extract_requirements_with_report(FIXTURES / "ecommerce_srs_ieee.docx", VISION)
        assert [(r["id"], r["text"]) for r in vision] == [(r["id"], r["text"]) for r in local]
        assert len(client.calls) == 1 and report.figures[0]["caption"] == "Figure 1: System context diagram"

    def test_unknown_diagram_type_is_other(self):
        reading = VisionEngine(TAGS, client=None).parse({"diagram_type": "mind map", "requirements": "nonsense"})
        assert reading.diagram_type == "other" and reading.statements == ()


# ---------------------------------------------------------------------------
# Options / config
# ---------------------------------------------------------------------------

class TestOptions:
    def test_env_default_and_validation(self, monkeypatch):
        monkeypatch.setenv("SRS_DIAGRAM_MODE", "vision")
        assert ExtractionOptions().effective_diagram_mode == "vision"
        assert ExtractionOptions(diagram_mode="off").effective_diagram_mode == "off"
        monkeypatch.setenv("SRS_DIAGRAM_MODE", "bogus")
        assert ExtractionOptions().effective_diagram_mode == "local"
        with pytest.raises(ValueError):
            ExtractionOptions(diagram_mode="bogus")

    def test_from_mapping_and_report(self):
        opts = ExtractionOptions.from_mapping({"diagram_mode": "Vision", "max_figures": 5})
        assert (opts.diagram_mode, opts.max_figures) == ("vision", 5)
        assert opts.to_dict()["diagram_mode"] == "vision"

    def test_traceability_yml(self):
        cfg = config_from_dict({"version": 1, "srs": "docs/SRS.docx", "openapi": "openapi.yaml",
                                "srs_extraction": {"diagram_mode": "vision", "max_figures": 3}})
        opts = cfg.extraction_options()
        assert (opts.diagram_mode, opts.max_figures) == ("vision", 3)
        with pytest.raises(ConfigError):
            config_from_dict({"version": 1, "srs": "docs/SRS.docx", "openapi": "openapi.yaml",
                              "srs_extraction": {"diagram_mode": "ocr-please"}})

    def test_report_round_trip_keeps_figures(self, tmp_path):
        path = _docx(tmp_path, [("h1", "Requirements"), REQ_TABLE,
                                ("img", "REQ-09: The system shall require 2FA for orders above $500.", None)])
        _, report = extract_requirements_with_report(path, LOCAL)
        from src.srs import ExtractionReport

        again = ExtractionReport.from_dict(json.loads(json.dumps(report.to_dict())))
        assert again.figures == report.figures and "1/1 figure(s) read" in again.summary()


class TestReviewProvenance:
    def test_violation_names_the_figure(self):
        from src.review.engine import requirement_source

        req = {"id": "REQ-09", "strategy": "diagram", "figures": [
            {"figure": "FIG-2", "caption": "Figure 2: REQ-09 2FA", "diagram_type": "activity",
             "engine": "vision:gpt-4o-mini"}]}
        assert requirement_source(req) == "SRS Figure 2: REQ-09 2FA · activity diagram · read by vision:gpt-4o-mini"
        assert requirement_source({"id": "REQ-04", "strategy": "table"}) == ""
