"""
scripts/e2e_system_test.py — the whole product lifecycle, end to end
=====================================================================
Drives the real FastAPI app, the real SRS parser, the real embedding model and
the real GitHub App review flow, in one throw-away sandbox:

  1. UPLOAD    build an IEEE-830 style SRS .docx in which one requirement
               (REQ-09, two-factor authentication) exists ONLY as a UML activity
               diagram image, upload it with the base and v2 OpenAPI specs
               through POST /api/projects/{id}/upload, and check the parser read
               the diagram (and skipped the appendix image).
  2. ANALYSIS  POST /analyze, then the matrix, drift and benchmark endpoints:
               the diagram requirement must trace to the 2FA/checkout endpoints
               and its drift in API v2 must be detected.
  3. PR BOT    a simulated GitHub (HTTP layer only) hosts a repo with the SRS,
               the spec and .traceability.yml. Four signed pull_request webhooks
               go to POST /api/github/webhook; the real ReviewFlow reviews each
               PR and the check run / PR comment / stored review are checked:
                 PASS  additive change (optional discount_code)
                 BLOCK checkout loses 2FA       -> REQ-09, the diagram requirement
                 BLOCK refund loses transaction_id -> REQ-04
                 WARN  wishlist PR delivers only part of its linked issue

Nothing touches the developer's database, upload folder, Qdrant collection or
GitHub: the DB / uploads live in a temp dir, Qdrant is bypassed (NumPy
similarity) and GitHub is an httpx.MockTransport.

The diagram is read by the vision model. Offline (default) the model's answer
is replayed from tests/fixtures/e2e/vision_recordings.json, a recording of a
real gpt-4o-mini reading of the same diagram; --live-vision calls the model.

Usage
    python scripts/e2e_system_test.py                     # offline, ~1-2 min (loads mpnet)
    python scripts/e2e_system_test.py --live-vision       # real vision model (OPENAI_API_KEY)
    python scripts/e2e_system_test.py --live-vision --record   # refresh the recording
    python scripts/e2e_system_test.py --live-llm          # real LLM explanations in PR comments
    python scripts/e2e_system_test.py --keep out/e2e      # keep SRS, specs, comments, report

Exit code 0 = every check passed, 1 = at least one failed. tests/test_e2e_lifecycle.py
runs the offline variant under pytest (marked slow).
"""

from __future__ import annotations

import argparse
import base64
import copy
import hashlib
import hmac
import io
import json
import os
import re
import sys
import tempfile
import time
from dataclasses import dataclass, field
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Callable, Dict, List, Optional, Tuple
from urllib.parse import unquote

ROOT = Path(__file__).resolve().parent.parent
for p in (ROOT, ROOT / "backend"):
    if str(p) not in sys.path:
        sys.path.insert(0, str(p))

RECORDINGS = ROOT / "tests" / "fixtures" / "e2e" / "vision_recordings.json"
STORE_V1 = ROOT / "data" / "api_specs" / "store_v1.json"
STORE_V2 = ROOT / "data" / "api_specs" / "store_v2.json"
MD_SRS = ROOT / "data" / "requirements" / "ecommerce_srs.md"
GT_REQ09 = {"POST /api/v1/checkout", "POST /api/v1/auth/2fa/challenge", "POST /api/v1/auth/2fa/verify"}

OWNER, REPO = "acme", "nexus-store"
APP_ID, INSTALLATION, SECRET = "424242", 77, "e2e-webhook-secret"
BASE_SHA = "b" * 40


# ===========================================================================
# Check log
# ===========================================================================

@dataclass
class Check:
    stage: str
    name: str
    ok: bool
    detail: str = ""


@dataclass
class Checks:
    rows: List[Check] = field(default_factory=list)
    quiet: bool = False

    def __call__(self, stage: str, name: str, ok: bool, detail: Any = "") -> bool:
        row = Check(stage, name, bool(ok), str(detail))
        self.rows.append(row)
        if not self.quiet:
            mark = "PASS" if row.ok else "FAIL"
            print(f"  [{mark}] {name}" + (f"  ({row.detail})" if row.detail else ""))
        return row.ok

    @property
    def failed(self) -> List[Check]:
        return [r for r in self.rows if not r.ok]

    def markdown(self) -> str:
        lines = ["| Stage | Check | Result | Detail |", "|---|---|---|---|"]
        for r in self.rows:
            lines.append(f"| {r.stage} | {r.name} | {'PASS' if r.ok else '**FAIL**'} | {r.detail.replace('|', '/')} |")
        return "\n".join(lines)


def stage(title: str) -> None:
    print(f"\n=== {title} " + "=" * max(0, 66 - len(title)))


# ===========================================================================
# 1. The industrial SRS, with a UML activity diagram
# ===========================================================================

def _font(size: int, bold: bool = False):
    from PIL import ImageFont

    for name in (("arialbd.ttf", "DejaVuSans-Bold.ttf") if bold else ("arial.ttf", "DejaVuSans.ttf")):
        try:
            return ImageFont.truetype(name, size)
        except OSError:
            continue
    return ImageFont.load_default()


def activity_diagram_png() -> bytes:
    """UML activity diagram for REQ-09: the requirement is in the picture only."""
    from PIL import Image, ImageDraw

    W, H = 1200, 1500
    img = Image.new("RGB", (W, H), "white")
    d = ImageDraw.Draw(img)
    title, body, small = _font(30, True), _font(22), _font(19)
    ink = "black"

    def centered(box, text, font):
        x0, y0, x1, y1 = box
        lines = text.split("\n")
        heights = [d.textbbox((0, 0), ln, font=font)[3] for ln in lines]
        y = (y0 + y1) / 2 - (sum(heights) + 6 * (len(lines) - 1)) / 2
        for ln, h in zip(lines, heights):
            w = d.textlength(ln, font=font)
            d.text(((x0 + x1) / 2 - w / 2, y), ln, fill=ink, font=font)
            y += h + 6

    def action(cx, cy, text, w=620, h=96):
        box = (cx - w / 2, cy - h / 2, cx + w / 2, cy + h / 2)
        d.rounded_rectangle(box, radius=28, outline=ink, width=3)
        centered(box, text, body)
        return box

    def decision(cx, cy, w=200, h=110):
        d.polygon([(cx, cy - h / 2), (cx + w / 2, cy), (cx, cy + h / 2), (cx - w / 2, cy)], outline=ink, width=3)

    def arrow(p0, p1, label=""):
        d.line([p0, p1], fill=ink, width=3)
        (x0, y0), (x1, y1) = p0, p1
        if abs(x1 - x0) > abs(y1 - y0):
            s = 1 if x1 > x0 else -1
            d.polygon([(x1, y1), (x1 - 16 * s, y1 - 9), (x1 - 16 * s, y1 + 9)], fill=ink)
        else:
            s = 1 if y1 > y0 else -1
            d.polygon([(x1, y1), (x1 - 9, y1 - 16 * s), (x1 + 9, y1 - 16 * s)], fill=ink)
        if label:
            d.text(((x0 + x1) / 2 + 10, (y0 + y1) / 2 - 26), label, fill=ink, font=small)

    d.text((40, 30), "REQ-09  Two-Factor Authentication Enforcement", fill=ink, font=title)
    d.text((40, 74), "UML activity diagram - high-value operations", fill=ink, font=body)
    d.line([(40, 112), (W - 40, 112)], fill=ink, width=2)

    cx = 470
    d.ellipse((cx - 20, 140, cx + 20, 180), fill=ink)                               # initial node
    arrow((cx, 180), (cx, 222))
    action(cx, 270, "Customer submits checkout or\nregisters a new payment method")
    arrow((cx, 318), (cx, 372))
    decision(cx, 427)
    centered((cx - 90, 400, cx + 90, 455), "high\nvalue?", small)
    d.text((cx + 120, 360), "[order total > $500 OR\n new payment method]", fill=ink, font=small)
    # [no] branch
    arrow((cx + 100, 427), (cx + 330, 427), "[no]")
    action(cx + 470, 427, "Process operation\nnormally", w=260, h=90)
    arrow((cx + 470, 472), (cx + 470, 1342))
    # [yes] branch
    arrow((cx, 482), (cx, 552), "[yes]")
    action(cx, 600, "POST /api/v1/auth/2fa/challenge\nissue time-limited OTP (expires in 5 min)")
    arrow((cx, 648), (cx, 702))
    action(cx, 750, "Customer enters the OTP code")
    arrow((cx, 798), (cx, 852))
    action(cx, 900, "POST /api/v1/auth/2fa/verify\nvalidate the submitted code")
    arrow((cx, 948), (cx, 1002))
    decision(cx, 1057)
    centered((cx - 90, 1030, cx + 90, 1085), "code\nvalid?", small)
    arrow((cx - 100, 1057), (cx - 290, 1057), "[no]")
    action(cx - 300, 1180, "Reject operation\nHTTP 403 two_fa_required", w=300, h=96)
    arrow((cx - 300, 1057), (cx - 300, 1132))
    arrow((cx, 1112), (cx, 1182), "[yes]")
    action(cx, 1230, "Authorise: POST /api/v1/checkout\nwith the verified two_fa_token")
    arrow((cx, 1278), (cx, 1342))
    d.ellipse((cx - 24, 1342, cx + 24, 1390), outline=ink, width=3)                   # final node
    d.ellipse((cx - 13, 1353, cx + 13, 1379), fill=ink)
    d.line([(cx + 24, 1366), (cx + 470, 1366)], fill=ink, width=3)

    note = (W - 330, 560, W - 30, 800)                                              # UML note
    d.polygon([(note[0], note[1]), (note[2] - 30, note[1]), (note[2], note[1] + 30), (note[2], note[3]),
               (note[0], note[3])], outline=ink, width=2)
    d.text((note[0] + 16, note[1] + 16),
           "<<constraint>>\n2FA must be verified\nBEFORE the high-value\noperation is authorised.\n"
           "Applies to checkout\nand to adding a\npayment method.", fill=ink, font=small, spacing=6)

    buf = io.BytesIO()
    img.save(buf, format="PNG")
    return buf.getvalue()


def context_diagram_png() -> bytes:
    """System context diagram: architecture, no behaviour of its own."""
    from PIL import Image, ImageDraw

    img = Image.new("RGB", (900, 360), "white")
    d = ImageDraw.Draw(img)
    f = _font(24)
    for x, label in ((40, "Shopper\n(web / mobile)"), (340, "Nexus Store API"), (640, "Payment\nGateway")):
        d.rectangle((x, 120, x + 220, 240), outline="black", width=3)
        d.multiline_text((x + 20, 150), label, fill="black", font=f, spacing=6)
    for x in (260, 560):
        d.line([(x, 180), (x + 80, 180)], fill="black", width=3)
        d.polygon([(x + 80, 180), (x + 64, 171), (x + 64, 189)], fill="black")
    d.text((40, 30), "Nexus Store - system context", fill="black", font=_font(28, True))
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    return buf.getvalue()


def mockup_png() -> bytes:
    from PIL import Image, ImageDraw

    img = Image.new("RGB", (640, 400), (245, 245, 245))
    d = ImageDraw.Draw(img)
    d.rectangle((20, 20, 620, 80), outline="black", width=2)
    d.text((40, 40), "REQ-77: The system shall show a dark-mode toggle (mock-up only)", fill="black", font=_font(18))
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    return buf.getvalue()


def _requirement_bodies() -> Dict[str, Tuple[str, str]]:
    """{id: (title, body)} from the plain-text demo SRS, so scores stay comparable."""
    from src.srs.readers.markdown_reader import legacy_extract

    out = {}
    for r in legacy_extract(MD_SRS.read_text(encoding="utf-8")):
        prefix = f"{r['title']}. "
        out[r["id"]] = (r["title"], r["text"][len(prefix):] if r["text"].startswith(prefix) else r["text"])
    return out


def build_srs_docx(path: Path) -> Path:
    """IEEE-830 layout: cover, revision history, TOC, definitions, a context
    diagram, a requirement table (REQ-01..06), tagged headings (REQ-07, 08, 10),
    REQ-09 ONLY as a UML activity diagram, and an appendix mock-up."""
    from docx import Document
    from docx.shared import Inches

    reqs = _requirement_bodies()
    doc = Document()
    doc.add_heading("Software Requirements Specification: Nexus Store Platform", level=0)
    doc.add_paragraph("Version 3.0 - Approved baseline - Prepared by the Nexus platform team")

    doc.add_heading("Revision History", level=1)
    t = doc.add_table(rows=3, cols=3)
    for r, row in enumerate([("Version", "Date", "Change"), ("2.0", "2026-03-01", "Payments chapter"),
                             ("3.0", "2026-10-01", "Security workflows as UML")]):
        for c, val in enumerate(row):
            t.cell(r, c).text = val
    doc.add_heading("Table of Contents", level=1)
    for line in ("1 Introduction ........ 2", "2 Overall Description ........ 3", "3 Specific Requirements ........ 4"):
        doc.add_paragraph(line)

    doc.add_heading("1 Introduction", level=1)
    doc.add_heading("1.1 Purpose", level=2)
    doc.add_paragraph("This SRS specifies the externally visible behaviour of the Nexus Store API. Every API "
                      "change is reviewed against it before merge.")
    doc.add_heading("1.2 Definitions, Acronyms and Abbreviations", level=2)
    g = doc.add_table(rows=3, cols=2)
    for r, row in enumerate([("Term", "Definition"), ("OTP", "One-time password"), ("2FA", "Two-factor authentication")]):
        for c, val in enumerate(row):
            g.cell(r, c).text = val

    doc.add_heading("2 Overall Description", level=1)
    doc.add_paragraph("Shoppers use the web and mobile apps; the Nexus Store API talks to the payment gateway.")
    doc.add_paragraph().add_run().add_picture(io.BytesIO(context_diagram_png()), width=Inches(5.5))
    doc.add_paragraph("Figure 1: System context diagram", style="Caption")

    doc.add_heading("3 Specific Requirements", level=1)
    doc.add_heading("3.1 Core Functional Requirements", level=2)
    table = doc.add_table(rows=1, cols=4)
    for c, h in enumerate(("Req ID", "Title", "Description", "Priority")):
        table.rows[0].cells[c].text = h
    for rid in ("REQ-01", "REQ-02", "REQ-03", "REQ-04", "REQ-05", "REQ-06"):
        cells = table.add_row().cells
        title, body = reqs[rid]
        for c, val in enumerate((rid, title, body, "High")):
            cells[c].text = val

    doc.add_heading("3.2 Account and Order Features", level=2)
    for rid in ("REQ-07", "REQ-08", "REQ-10"):
        title, body = reqs[rid]
        doc.add_heading(f"{rid} {title}", level=3)
        doc.add_paragraph(body)

    doc.add_heading("3.3 Security Workflows", level=2)
    doc.add_paragraph("High-value operations follow the workflow modelled in Figure 2.")
    doc.add_paragraph().add_run().add_picture(io.BytesIO(activity_diagram_png()), width=Inches(6))
    doc.add_paragraph("Figure 2: REQ-09 Two-Factor Authentication Enforcement - activity diagram", style="Caption")

    doc.add_heading("Appendix A: UI Mock-ups", level=1)
    doc.add_paragraph().add_run().add_picture(io.BytesIO(mockup_png()), width=Inches(4))
    doc.add_paragraph("Figure 3: Settings screen mock-up", style="Caption")

    doc.save(path)
    return path


# ===========================================================================
# Vision model: live, recorded, counted
# ===========================================================================

class CountingVision:
    """Wraps a vision client; counts calls and keeps each answer (for --record)."""

    def __init__(self, inner: Any):
        self.inner = inner
        self.calls: List[Dict[str, Any]] = []
        self.answers: Dict[str, Dict[str, Any]] = {}
        self.chat = self
        self.completions = self

    def create(self, **kwargs: Any) -> Any:
        self.calls.append(kwargs)
        response = self.inner.chat.completions.create(**kwargs)
        prompt = kwargs["messages"][1]["content"][0]["text"]
        self.answers[_figure_key(prompt)] = json.loads(response.choices[0].message.content or "{}")
        return response


class ReplayVision:
    """Answers with the recorded model output for the figure named in the prompt."""

    def __init__(self, recordings: Dict[str, Any]):
        self.recordings = recordings
        self.chat = self
        self.completions = self

    def create(self, **kwargs: Any) -> Any:
        from types import SimpleNamespace

        prompt = kwargs["messages"][1]["content"][0]["text"]
        answer = self.recordings.get(_figure_key(prompt), {"diagram_type": "other", "requirements": []})
        return SimpleNamespace(choices=[SimpleNamespace(message=SimpleNamespace(content=json.dumps(answer)))])


def _figure_key(prompt: str) -> str:
    m = re.search(r"Figure caption: (Figure \d+)", prompt)
    return m.group(1) if m else "unknown"


# ===========================================================================
# 3. A simulated GitHub (HTTP layer only)
# ===========================================================================

def _iso(ts: float) -> str:
    return datetime.fromtimestamp(ts, tz=timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")


class SimulatedGitHub:
    """Enough of api.github.com for ReviewFlow: installation tokens, contents at
    a ref, PR files, compare, issues, GraphQL closing issues, check runs and
    issue comments. Files are stored per commit SHA."""

    def __init__(self) -> None:
        self.files: Dict[str, Dict[str, bytes]] = {BASE_SHA: {}}
        self.pr_files: Dict[int, List[Dict[str, Any]]] = {}
        self.prs: Dict[int, Dict[str, Any]] = {}
        self.issues: Dict[int, Dict[str, Any]] = {}
        self.closing: Dict[int, List[int]] = {}
        self.checks: Dict[int, Dict[str, Any]] = {}
        self.comments: Dict[int, List[Dict[str, Any]]] = {}
        self._next_comment = 1000

    def put(self, sha: str, path: str, data: Any) -> None:
        if isinstance(data, (dict, list)):
            data = json.dumps(data, indent=1)
        self.files.setdefault(sha, {})[path] = data if isinstance(data, bytes) else data.encode()

    def __call__(self, request: Any) -> Any:
        import httpx

        method, path = request.method, unquote(request.url.path)
        body = json.loads(request.content) if request.content else {}
        if path.endswith("/access_tokens"):
            return httpx.Response(201, json={"token": "ghs_e2e", "expires_at": _iso(time.time() + 3600)})
        if path == "/graphql":
            number = body.get("variables", {}).get("number") or body.get("variables", {}).get("pr")
            nodes = []
            for n in self.closing.get(int(number or 0), []):
                i = self.issues[n]
                nodes.append({"number": n, "title": i["title"], "body": i["body"], "url": i["html_url"],
                              "repository": {"nameWithOwner": f"{OWNER}/{REPO}"}})
            return httpx.Response(200, json={"data": {"repository": {"pullRequest": {
                "closingIssuesReferences": {"nodes": nodes}}}}})
        prefix = f"/repos/{OWNER}/{REPO}/"
        if not path.startswith(prefix):
            return httpx.Response(404, json={"message": "Not Found"})
        rest = path[len(prefix):]

        if method == "POST" and rest == "check-runs":
            cid = len(self.checks) + 1
            self.checks[cid] = {"id": cid, "html_url": f"https://github.com/{OWNER}/{REPO}/runs/{cid}", **body}
            return httpx.Response(201, json=self.checks[cid])
        if m := re.fullmatch(r"check-runs/(\d+)", rest):
            self.checks[int(m.group(1))].update(body)
            return httpx.Response(200, json=self.checks[int(m.group(1))])
        if re.fullmatch(r"compare/\w+\.\.\.\w+", rest):
            return httpx.Response(200, json={"merge_base_commit": {"sha": BASE_SHA}})
        if m := re.fullmatch(r"pulls/(\d+)/files", rest):
            return httpx.Response(200, json=self.pr_files.get(int(m.group(1)), []))
        if m := re.fullmatch(r"pulls/(\d+)", rest):
            return httpx.Response(200, json=self.prs[int(m.group(1))])
        if rest.startswith("contents/"):
            data = self.files.get(request.url.params.get("ref"), {}).get(rest[len("contents/"):])
            if data is None:
                return httpx.Response(404, json={"message": "Not Found"})
            return httpx.Response(200, json={"type": "file", "encoding": "base64", "size": len(data),
                                             "content": base64.b64encode(data).decode()})
        if m := re.fullmatch(r"issues/(\d+)", rest):
            issue = self.issues.get(int(m.group(1)))
            return httpx.Response(200, json=issue) if issue else httpx.Response(404, json={"message": "Not Found"})
        if m := re.fullmatch(r"issues/(\d+)/comments", rest):
            thread = self.comments.setdefault(int(m.group(1)), [])
            if method == "GET":
                return httpx.Response(200, json=thread)
            self._next_comment += 1
            thread.append({"id": self._next_comment, "body": body["body"],
                           "html_url": f"https://github.com/{OWNER}/{REPO}/pull/{m.group(1)}#c{self._next_comment}"})
            return httpx.Response(201, json=thread[-1])
        if m := re.fullmatch(r"issues/comments/(\d+)", rest):
            for thread in self.comments.values():
                for c in thread:
                    if c["id"] == int(m.group(1)):
                        c["body"] = body["body"]
                        return httpx.Response(200, json=c)
        return httpx.Response(404, json={"message": f"unrouted {method} {rest}"})


@dataclass
class PRScenario:
    key: str
    number: int
    title: str
    issue_title: str
    issue_body: str
    mutate: Callable[[Dict[str, Any]], None]
    expected: str
    expected_regressed: List[str]


def _scenario_head(key: str) -> Dict[str, Any]:
    from src.pr_reviewer import SIMULATED_PR_SCENARIOS

    return copy.deepcopy(SIMULATED_PR_SCENARIOS[key]["changes"][0]["head"])


def _drop_2fa(spec: Dict[str, Any]) -> None:
    spec["paths"]["/api/v1/checkout"]["post"] = _scenario_head("drop_2fa_requirement")


def _add_wishlist(spec: Dict[str, Any]) -> None:
    spec["paths"]["/api/v1/wishlist/items"] = {"post": {
        "summary": "Save a product to the user's wishlist",
        "description": "Adds a product to the authenticated user's wishlist so it can be bought later. "
                       "Returns the wishlist with product_id and added_at for each saved product.",
        "requestBody": {"required": True, "content": {"application/json": {"schema": {
            "type": "object", "required": ["product_id"],
            "properties": {"product_id": {"type": "string", "description": "Product to save"}}}}}},
        "responses": {"201": {"description": "Product saved to the wishlist"}},
    }}


SCENARIOS = [
    PRScenario("discount", 101, "feat: optional discount_code on add-to-cart", "Support discount codes",
               "Shoppers want promo codes.\n\n- [ ] Adding an item to the cart accepts an optional discount code",
               lambda s: s["paths"]["/api/v1/cart/items"].__setitem__("post", _scenario_head("add_endpoint")),
               "PASS", []),
    PRScenario("frictionless", 102, "refactor: frictionless checkout (remove 2FA step-up)",
               "Reduce checkout friction", "Conversion drops at the 2FA step. Make checkout one click.",
               _drop_2fa, "BLOCK", ["REQ-09"]),
    PRScenario("refund", 103, "refactor: server-generated refund_reference", "Simplify refunds",
               "## Acceptance criteria\n- Refunds use a server-generated reference",
               lambda s: s["paths"]["/api/v1/refund"].__setitem__("post", _scenario_head("remove_idempotency")),
               "BLOCK", ["REQ-04"]),
    PRScenario("wishlist", 104, "feat: wishlist", "Wishlist",
               "- [ ] Users can save a product to a wishlist\n"
               "- [ ] Users receive an email when the price of a wishlisted product drops below a target price",
               _add_wishlist, "WARN", []),
]


# ===========================================================================
# The lifecycle
# ===========================================================================

def _rsa_pem() -> str:
    from cryptography.hazmat.primitives import serialization
    from cryptography.hazmat.primitives.asymmetric import rsa

    key = rsa.generate_private_key(public_exponent=65537, key_size=2048)
    return key.private_bytes(serialization.Encoding.PEM, serialization.PrivateFormat.PKCS8,
                             serialization.NoEncryption()).decode()


class _NoVectorStore:
    """Keeps the run off the developer's Qdrant (NumPy similarity instead)."""

    is_available = False


def run_lifecycle(workdir: Path, *, live_vision: bool = False, live_llm: bool = False, record: bool = False,
                  quiet: bool = False) -> Tuple[Checks, Dict[str, Any]]:
    """Run all three stages. Returns the checks and the artefacts (requirements,
    matrix rows, PR comments) for the report."""
    import httpx
    from fastapi.testclient import TestClient
    from sqlalchemy import create_engine
    from sqlalchemy.ext.asyncio import async_sessionmaker, create_async_engine
    from sqlalchemy.orm import Session
    from sqlalchemy.pool import NullPool

    from app.core.config import settings
    from app.core.security import get_current_user
    from app.database import Base, User
    from app.main import app
    from app.routers import github as gh_router
    from app.services import pipeline
    from app.services.github import GitHubAppAuth
    from app.services.github.review_flow import ReviewFlow
    from src.explainer import ComplianceExplainer
    from src.srs import diagrams

    checks = Checks(quiet=quiet)
    art: Dict[str, Any] = {}
    workdir.mkdir(parents=True, exist_ok=True)

    # ---- sandbox -------------------------------------------------------------
    if "sqlite" not in settings.DATABASE_URL or not str(settings.UPLOAD_DIR).startswith(tempfile.gettempdir()):
        raise RuntimeError("Refusing to run: the backend is not pointed at a temp DB / upload dir.")
    saved_env = {k: os.environ.get(k) for k in ("SRS_DIAGRAM_MODE", "SRS_DIAGRAM_CACHE_DIR")}
    os.environ["SRS_DIAGRAM_MODE"] = "vision"
    os.environ["SRS_DIAGRAM_CACHE_DIR"] = str(workdir / "diagram_cache")      # fresh: really reads the figures
    saved = {
        "vision": diagrams.vision_client_factory, "vs": pipeline.get_vector_store,
        "secret": settings.GITHUB_WEBHOOK_SECRET, "app_id": settings.GITHUB_APP_ID,
    }
    if live_vision:
        real = diagrams.default_vision_client()
        if real is None:
            raise RuntimeError("--live-vision needs OPENAI_API_KEY in the environment or .env")
        vision = CountingVision(real)
    else:
        vision = CountingVision(ReplayVision(json.loads(RECORDINGS.read_text(encoding="utf-8"))))
    diagrams.vision_client_factory = lambda: vision
    pipeline.get_vector_store = lambda: _NoVectorStore()
    settings.GITHUB_WEBHOOK_SECRET, settings.GITHUB_APP_ID = SECRET, APP_ID

    sync = create_engine(settings.DATABASE_URL.replace("+aiosqlite", ""))
    Base.metadata.create_all(sync)
    with Session(sync, expire_on_commit=False) as s:
        user = User(google_id=f"e2e-{time.time_ns()}", email="e2e@example.com", name="E2E")
        s.add(user)
        s.commit()
    app.dependency_overrides[get_current_user] = lambda: user
    client = TestClient(app)

    try:
        # ===== 1. UPLOAD =====================================================
        stage("1. Upload: industrial SRS with a UML diagram + OpenAPI specs")
        srs_path = build_srs_docx(workdir / "Nexus_Store_SRS_v3.docx")
        (workdir / "activity_diagram_REQ-09.png").write_bytes(activity_diagram_png())
        r = client.post("/api/projects", json={"name": "E2E Nexus Store", "description": "lifecycle test"})
        checks("upload", "create project (POST /api/projects)", r.status_code == 201, r.status_code)
        pid = r.json()["id"]

        def upload(doc_type: str, path: Path, mime: str) -> Any:
            return client.post(f"/api/projects/{pid}/upload", data={"doc_type": doc_type},
                               files={"file": (path.name, path.read_bytes(), mime)})

        t0 = time.time()
        r = upload("SRS", srs_path, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
        ext = r.json().get("extraction") or {}
        checks("upload", "SRS .docx accepted and parsed at upload", r.status_code == 201,
               f"{r.status_code}, {ext.get('n_requirements')} requirements, {ext.get('n_figures')} figures read, "
               f"{time.time() - t0:.1f}s")
        checks("upload", "all 10 requirements extracted (REQ-01..REQ-10), nothing extra",
               sorted(ext.get("requirement_ids") or []) == [f"REQ-{i:02d}" for i in range(1, 11)],
               ext.get("requirement_ids"))
        for doc_type, path in (("API_V1", STORE_V1), ("API_V2", STORE_V2)):
            r = upload(doc_type, path, "application/json")
            checks("upload", f"{doc_type} spec accepted ({path.name})", r.status_code == 201, r.status_code)

        preview = client.get(f"/api/projects/{pid}/srs/requirements").json()
        reqs = {q["id"]: q for q in preview["requirements"]}
        report = preview["report"]
        art["requirements"] = preview["requirements"]
        art["extraction_report"] = report
        req09 = reqs.get("REQ-09", {})
        checks("upload", "REQ-09 was read from the UML activity diagram",
               req09.get("strategy") == "diagram" and req09.get("source", {}).get("locator") == "figure 2",
               f"strategy={req09.get('strategy')}, locator={req09.get('source', {}).get('locator')}")
        text09 = req09.get("text", "").lower()
        checks("upload", "diagram requirement keeps the diagram's facts (2FA, $500, OTP endpoints)",
               "500" in text09 and ("2fa" in text09 or "two-factor" in text09) and "otp" in text09,
               req09.get("text", "")[:160] + "...")
        by_fig = {f["figure"]: f for f in report.get("figures", [])}
        checks("upload", "context diagram read but adds no requirement (structural = context only)",
               "FIG-1" in by_fig and by_fig["FIG-1"]["requirements"] == [],
               f"{by_fig.get('FIG-1', {}).get('diagram_type')}: {by_fig.get('FIG-1', {}).get('summary', '')[:90]}")
        checks("upload", "appendix mock-up never sent to the vision model",
               len(vision.calls) == 2 and "REQ-77" not in reqs, f"{len(vision.calls)} vision call(s)")
        checks("upload", "boilerplate removed (revision history, TOC, glossary)",
               not any(re.search(r"Revision|Table of Contents|One-time password", q["text"]) for q in reqs.values())
               and any(k.startswith("blocklisted_section") for k in report["skipped"]),
               ", ".join(k for k in report["skipped"] if k.startswith("blocklisted")))
        checks("upload", "strategies: table + tagged headings + diagram",
               report["by_strategy"] == {"table": 6, "tagged": 3, "diagram": 1}, report["by_strategy"])

        # ===== 2. ANALYSIS ===================================================
        stage("2. Analysis: traceability matrix, drift, benchmark")
        t0 = time.time()
        r = client.post(f"/api/projects/{pid}/analyze", json={"enable_drift": True})
        body = r.json()
        checks("analysis", "POST /analyze succeeded", r.status_code == 200,
               f"{body.get('n_requirements')} reqs x {body.get('n_endpoints')} endpoints, "
               f"compliance {body.get('overall_compliance')}, {time.time() - t0:.1f}s")
        checks("analysis", "vision cache made the re-parse free and identical", len(vision.calls) == 2,
               f"{len(vision.calls)} vision call(s) in total")
        matrix = client.get(f"/api/projects/{pid}/matrix").json()
        art["matrix"] = matrix
        top09 = _top_endpoints(matrix, "REQ-09")
        checks("analysis", "REQ-09 (from the diagram) traces to the 2FA / checkout endpoints",
               bool(top09) and top09[0][0] in GT_REQ09,
               ", ".join(f"{e} {s:.3f}" for e, s in top09[:3]))
        drift = client.get(f"/api/projects/{pid}/drift").json()
        art["drift"] = drift
        drifted = _drifted_requirements(drift)
        checks("analysis", "API v2 drift on REQ-09 detected", "REQ-09" in drifted, sorted(drifted))
        bench = client.get(f"/api/projects/{pid}/benchmark").json()
        art["benchmark"] = bench
        rep = bench.get("report", {})
        f1_t = _metric(rep, "Transformer", "F1-Score")
        f1_b = _metric(rep, "TF-IDF", "F1-Score")
        checks("analysis", "semantic model beats the TF-IDF baseline (F1)",
               f1_t is not None and f1_b is not None and f1_t > f1_b,
               f"F1 {f1_t} vs {f1_b}, MAP {bench.get('transformer_map')} vs {bench.get('tfidf_map')}")

        # ===== 3. PR BOT =====================================================
        stage("3. PR bot: signed webhooks -> GitHub App review flow")
        gh = SimulatedGitHub()
        config = {"version": 1, "srs": "docs/SRS.docx", "openapi": "openapi.json", "fail_on": "block",
                  "srs_extraction": {"diagram_mode": "vision"}}
        base_spec = json.loads(STORE_V1.read_text(encoding="utf-8"))
        for path, data in ((".traceability.yml", json.dumps(config)), ("docs/SRS.docx", srs_path.read_bytes()),
                           ("openapi.json", base_spec)):
            gh.put(BASE_SHA, path, data)

        http = httpx.AsyncClient(transport=httpx.MockTransport(gh))
        explainer = ComplianceExplainer()
        if not live_llm:
            explainer._available = False                     # deterministic template explanations
        flow = ReviewFlow(
            GitHubAppAuth(APP_ID, _rsa_pem(), http=http), http=http,
            embedder_factory=pipeline.get_embedder, explainer_factory=lambda: explainer,
            session_factory=async_sessionmaker(create_async_engine(settings.DATABASE_URL, poolclass=NullPool),
                                               expire_on_commit=False),
        )
        app.dependency_overrides[gh_router.get_review_handler] = lambda: flow
        gh_router.deliveries.clear()
        art["pr_comments"] = {}

        for sc in SCENARIOS:
            head_sha = hashlib.sha1(sc.key.encode()).hexdigest()
            head_spec = copy.deepcopy(base_spec)
            sc.mutate(head_spec)
            for path in (".traceability.yml", "docs/SRS.docx"):
                gh.put(head_sha, path, gh.files[BASE_SHA][path])
            gh.put(head_sha, "openapi.json", head_spec)
            gh.pr_files[sc.number] = [{"filename": "openapi.json", "status": "modified"}]
            issue_no = sc.number + 500
            gh.issues[issue_no] = {"number": issue_no, "title": sc.issue_title, "body": sc.issue_body,
                                   "html_url": f"https://github.com/{OWNER}/{REPO}/issues/{issue_no}"}
            gh.closing[sc.number] = [issue_no]
            pr = {"number": sc.number, "state": "open", "title": sc.title, "body": f"Closes #{issue_no}",
                  "html_url": f"https://github.com/{OWNER}/{REPO}/pull/{sc.number}", "user": {"login": "dev-e2e"},
                  "head": {"sha": head_sha, "ref": f"feature/{sc.key}"}, "base": {"sha": BASE_SHA, "ref": "main"}}
            gh.prs[sc.number] = pr
            payload = json.dumps({"action": "opened", "installation": {"id": INSTALLATION},
                                  "repository": {"name": REPO, "owner": {"login": OWNER}},
                                  "pull_request": pr}).encode()
            sig = "sha256=" + hmac.new(SECRET.encode(), payload, hashlib.sha256).hexdigest()
            n_checks = len(gh.checks)
            t0 = time.time()
            r = client.post("/api/github/webhook", content=payload, headers={
                "X-GitHub-Event": "pull_request", "X-GitHub-Delivery": f"e2e-{sc.key}",
                "X-Hub-Signature-256": sig, "Content-Type": "application/json"})
            run = gh.checks.get(n_checks + 1, {})
            output = run.get("output") or {}
            verdict = (output.get("title") or "").split(":")[0].strip()
            comment = (gh.comments.get(sc.number) or [{}])[-1].get("body", "")
            art["pr_comments"][sc.key] = comment
            regressed = sorted(set(re.findall(r"\| `(REQ-\d+|FIG-[\d.]+)` \|", output.get("summary") or "")))
            ok = (r.status_code == 202 and run.get("status") == "completed" and verdict == sc.expected
                  and regressed == sc.expected_regressed)
            checks("pr-bot", f"PR #{sc.number} {sc.title!r} -> {sc.expected}", ok,
                   f"webhook {r.status_code}, check {run.get('conclusion')}, title {output.get('title')!r}, "
                   f"regressed {regressed}, {time.time() - t0:.1f}s")
            if sc.expected == "BLOCK":
                checks("pr-bot", f"PR #{sc.number} comment explains the violation",
                       all(q in comment for q in sc.expected_regressed) and "<!-- traceability-bot -->" in comment,
                       f"{len(comment)} chars")
            if "REQ-09" in sc.expected_regressed:
                source = re.search(r"\| \*\*Defined in\*\* \| (.+?) \|", comment)
                checks("pr-bot", f"PR #{sc.number} comment points to the diagram that defines REQ-09",
                       bool(source) and "Figure 2" in source.group(1), source.group(1) if source else "missing")

        rows = client.get("/api/github/reviews").json()["reviews"]
        verdicts = sorted((row.get("pr_number"), row.get("verdict")) for row in rows)
        checks("pr-bot", "every review stored for the Live reviews page",
               verdicts == sorted((sc.number, sc.expected) for sc in SCENARIOS), verdicts)
    finally:
        for key, value in saved_env.items():
            if value is None:
                os.environ.pop(key, None)
            else:
                os.environ[key] = value
        diagrams.vision_client_factory = saved["vision"]
        pipeline.get_vector_store = saved["vs"]
        settings.GITHUB_WEBHOOK_SECRET, settings.GITHUB_APP_ID = saved["secret"], saved["app_id"]
        app.dependency_overrides.pop(get_current_user, None)
        app.dependency_overrides.pop(gh_router.get_review_handler, None)
        sync.dispose()

    if record and live_vision:
        RECORDINGS.parent.mkdir(parents=True, exist_ok=True)
        RECORDINGS.write_text(json.dumps({k: vision.answers[k] for k in sorted(vision.answers)}, indent=2) + "\n",
                              encoding="utf-8")
        print(f"\nRecorded {len(vision.answers)} vision answer(s) to {RECORDINGS.relative_to(ROOT)}")
    art["vision_answers"] = vision.answers
    return checks, art


# -- payload helpers (tolerant of the exact JSON shapes) -------------------------

def _top_endpoints(matrix: Dict[str, Any], req_id: str) -> List[Tuple[str, float]]:
    """[(endpoint id, score)] for one requirement, best first."""
    cells = [c for c in matrix.get("cells", []) if c.get("req_id") == req_id]
    return sorted(((c["ep_id"], float(c["score"])) for c in cells), key=lambda x: -x[1])


def _drifted_requirements(drift: Dict[str, Any]) -> set:
    """`pairs` already holds only the drifted requirement/endpoint pairs."""
    return {row.get("requirement_id") or row.get("req_id") for row in drift.get("pairs", [])}


def _metric(report: Dict[str, Any], method: str, column: str) -> Optional[float]:
    """report = EvaluationReport DataFrame .to_dict(): {column: {method label: value}}."""
    for label, value in (report.get(column) or {}).items():
        if label.lower().startswith(method.lower()):
            return round(float(value), 3)
    return None


# ===========================================================================
# CLI
# ===========================================================================

def main(argv: Optional[List[str]] = None) -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--live-vision", action="store_true", help="read the diagrams with the real vision model")
    ap.add_argument("--record", action="store_true", help="with --live-vision: save the answers as the recording")
    ap.add_argument("--live-llm", action="store_true", help="real LLM explanations in the PR comments")
    ap.add_argument("--keep", type=Path, help="keep the generated files and the report in this folder")
    args = ap.parse_args(argv)

    sandbox = Path(tempfile.mkdtemp(prefix="trace_e2e_"))
    os.environ["DATABASE_URL"] = f"sqlite+aiosqlite:///{(sandbox / 'e2e.db').as_posix()}"
    os.environ["UPLOAD_DIR"] = str(sandbox / "uploads")
    workdir = args.keep.resolve() if args.keep else sandbox / "work"

    started = time.time()
    checks, art = run_lifecycle(workdir, live_vision=args.live_vision, live_llm=args.live_llm, record=args.record)
    elapsed = time.time() - started

    (workdir / "e2e_report.md").write_text(
        f"# E2E lifecycle report\n\nvision: {'live' if args.live_vision else 'recorded'}, "
        f"LLM explanations: {'live' if args.live_llm else 'template'}, {elapsed:.0f}s\n\n{checks.markdown()}\n",
        encoding="utf-8")
    (workdir / "e2e_artifacts.json").write_text(json.dumps(
        {k: art.get(k) for k in ("requirements", "extraction_report", "drift", "benchmark", "vision_answers")},
        indent=1, default=str), encoding="utf-8")
    for key, comment in art.get("pr_comments", {}).items():
        (workdir / f"pr_comment_{key}.md").write_text(comment, encoding="utf-8")

    n_fail = len(checks.failed)
    print(f"\n{len(checks.rows) - n_fail}/{len(checks.rows)} checks passed in {elapsed:.0f}s. Files: {workdir}")
    return 1 if n_fail else 0


if __name__ == "__main__":
    sys.exit(main())
