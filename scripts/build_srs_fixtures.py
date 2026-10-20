"""
scripts/build_srs_fixtures.py — Generate the industrial SRS test fixtures
=========================================================================
Writes deterministic Word documents (and their golden *.expected.json) to
tests/fixtures/srs/. Each fixture reproduces a layout found in real SRS
documents, so the parser is tested against structure, not just plain text.

  ecommerce_srs_ieee.docx  IEEE 830 layout around the same 10 requirements as
                           data/requirements/ecommerce_srs.md: cover page,
                           revision table, TOC content control, boilerplate,
                           glossary table, image + caption, REQ-01..06 in a
                           requirement table, REQ-07..10 as tagged headings,
                           appendix with a decoy requirement.
  banking_srs.docx         BNK-REQ-01..08 as one key/value table each
                           (Requirement ID / Title / Description / Priority).
  messy_srs.docx           every edge case: FR1 / FR-2 / FR 3 / [REQ-4] /
                           FR-5.1 styles, merged + multi-paragraph cells,
                           repeated header row, vertical merge, 1x1 callout
                           table, content control, tracked insertion and
                           deletion, cross-reference, duplicate ID, bold
                           numbered pseudo-heading, NFR section.
  untagged_srs.docx        no IDs at all; only "shall/must" sentences.
  scanned.pdf              image-only PDF (no text layer).

The expected requirements for ecommerce/banking are taken from the Markdown
originals via the v1 parser, which is what makes the .docx -> .md parity test
meaningful.

Run:  python scripts/build_srs_fixtures.py
PDF:  python scripts/build_srs_fixtures.py --pdf   (Windows + Word: also exports
      ecommerce_srs_ieee.pdf through Word's "Save as PDF")
"""

from __future__ import annotations

import argparse
import io
import json
import subprocess
import sys
from pathlib import Path
from typing import Dict, List, Optional, Sequence

from docx import Document
from docx.document import Document as DocxDocument
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from src.srs.readers.markdown_reader import legacy_extract, read_text_file  # noqa: E402

OUT = ROOT / "tests" / "fixtures" / "srs"
ECOMMERCE_MD = ROOT / "data" / "requirements" / "ecommerce_srs.md"
BANKING_MD = ROOT / "data" / "samples" / "banking" / "aurora_banking_srs.md"

FIXED_DATE = "2026-10-01T09:00:00Z"


# ---------------------------------------------------------------------------
# python-docx helpers for things the public API does not expose
# ---------------------------------------------------------------------------

def _new_document() -> DocxDocument:
    doc = Document()
    styles = doc.styles
    for name in ("TOC 1", "TOC 2"):
        if name not in [s.name for s in styles]:
            styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    # deterministic core properties
    cp = doc.core_properties
    cp.author = "SRS fixture generator"
    cp.title = "Software Requirements Specification"
    import datetime as _dt
    stamp = _dt.datetime(2026, 10, 1, 9, 0, 0)
    cp.created = cp.modified = stamp
    cp.last_modified_by = "SRS fixture generator"
    cp.revision = 1
    return doc


def _header_footer(doc: DocxDocument, header: str, footer: str) -> None:
    section = doc.sections[0]
    section.header.paragraphs[0].text = header
    section.footer.paragraphs[0].text = footer


def _mark_header_row(table: Table, n: int = 1) -> None:
    for row in table.rows[:n]:
        tr_pr = row._tr.get_or_add_trPr()
        el = OxmlElement("w:tblHeader")
        tr_pr.append(el)


def _bold_row(table: Table, idx: int = 0) -> None:
    for cell in table.rows[idx].cells:
        for p in cell.paragraphs:
            for r in p.runs:
                r.bold = True


def _table(doc_or_cell, header: Sequence[str], rows: Sequence[Sequence[str]], style: str = "Table Grid") -> Table:
    table = doc_or_cell.add_table(rows=1 + len(rows), cols=len(header))
    try:
        table.style = style
    except (KeyError, ValueError):
        pass
    for c, text in enumerate(header):
        table.rows[0].cells[c].text = text
    for r, values in enumerate(rows, start=1):
        for c, text in enumerate(values):
            table.rows[r].cells[c].text = text
    _bold_row(table)
    return table


def _wrap_in_sdt(paragraphs: Sequence[Paragraph], gallery: Optional[str] = None, tag: str = "") -> None:
    """Move paragraphs into a block-level content control (w:sdt)."""
    first = paragraphs[0]._p
    sdt = OxmlElement("w:sdt")
    sdt_pr = OxmlElement("w:sdtPr")
    if tag:
        t = OxmlElement("w:tag")
        t.set(qn("w:val"), tag)
        sdt_pr.append(t)
    if gallery:
        obj = OxmlElement("w:docPartObj")
        g = OxmlElement("w:docPartGallery")
        g.set(qn("w:val"), gallery)
        u = OxmlElement("w:docPartUnique")
        obj.append(g)
        obj.append(u)
        sdt_pr.append(obj)
    content = OxmlElement("w:sdtContent")
    sdt.append(sdt_pr)
    sdt.append(content)
    first.addprevious(sdt)
    for p in paragraphs:
        content.append(p._p)


def _run(text: str, bold: bool = False):
    r = OxmlElement("w:r")
    if bold:
        rpr = OxmlElement("w:rPr")
        rpr.append(OxmlElement("w:b"))
        r.append(rpr)
    t = OxmlElement("w:t")
    t.set(qn("xml:space"), "preserve")
    t.text = text
    r.append(t)
    return r


def _tracked(kind: str, text: str, rev_id: int):
    """<w:ins> / <w:del> revision containing one run."""
    el = OxmlElement(f"w:{kind}")
    el.set(qn("w:id"), str(rev_id))
    el.set(qn("w:author"), "Reviewer")
    el.set(qn("w:date"), FIXED_DATE)
    r = OxmlElement("w:r")
    t = OxmlElement("w:delText" if kind == "del" else "w:t")
    t.set(qn("xml:space"), "preserve")
    t.text = text
    r.append(t)
    el.append(r)
    return el


def _png_bytes() -> bytes:
    from PIL import Image, ImageDraw

    img = Image.new("RGB", (320, 120), "white")
    d = ImageDraw.Draw(img)
    d.rectangle([10, 10, 150, 110], outline="black")
    d.rectangle([170, 10, 310, 110], outline="black")
    d.line([150, 60, 170, 60], fill="black")
    d.text((30, 50), "Client app", fill="black")
    d.text((190, 50), "Nexus API", fill="black")
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    return buf.getvalue()


def _toc(doc: DocxDocument, entries: Sequence[str]) -> None:
    paras = [doc.add_paragraph("Contents", style="TOC Heading")]
    for e in entries:
        paras.append(doc.add_paragraph(e, style="TOC 1"))
    _wrap_in_sdt(paras, gallery="Table of Contents")


def _write_expected(name: str, requirements: List[Dict[str, object]], extra: Optional[Dict] = None) -> None:
    payload = {"fixture": name, "requirements": requirements}
    if extra:
        payload.update(extra)
    (OUT / f"{Path(name).stem}.expected.json").write_text(
        json.dumps(payload, indent=2, ensure_ascii=False) + "\n", encoding="utf-8"
    )


def _md_requirements(path: Path) -> List[Dict[str, str]]:
    raw = read_text_file(path)
    out = []
    for r in legacy_extract(raw):
        body = r["text"][len(r["title"]) + 2:]           # text = "{title}. {body}"
        out.append({"id": r["id"], "title": r["title"], "text": r["text"], "body": body})
    return out


def _boilerplate_front(doc: DocxDocument, product: str) -> None:
    doc.add_heading("Software Requirements Specification", level=0)
    doc.add_paragraph(f"for {product}")
    doc.add_paragraph("Version 1.0 approved")
    doc.add_paragraph("Prepared by: Platform Engineering Team")
    doc.add_paragraph("2026-10-01")
    doc.add_page_break()

    doc.add_heading("Revision History", level=1)
    _table(doc, ["Version", "Date", "Author", "Description"], [
        ["0.1", "2026-08-12", "Analyst", "Initial draft of the requirement set"],
        ["1.0", "2026-10-01", "Architect", "Baseline approved by the change board"],
    ])


# ---------------------------------------------------------------------------
# Fixture 1 — IEEE 830 e-commerce SRS
# ---------------------------------------------------------------------------

def build_ecommerce() -> Path:
    reqs = _md_requirements(ECOMMERCE_MD)
    assert len(reqs) == 10, "data/requirements/ecommerce_srs.md must have 10 requirements"
    doc = _new_document()
    _header_footer(doc, "Nexus E-Commerce Platform SRS v1.0 - CONFIDENTIAL", "Page 1 of 9")
    _boilerplate_front(doc, "Nexus E-Commerce Platform")
    _toc(doc, ["1 Introduction\t3", "1.1 Purpose\t3", "2 Overall Description\t4",
               "3 Specific Requirements\t5", "Appendix A: Glossary\t9"])

    doc.add_heading("1 Introduction", level=1)
    doc.add_heading("1.1 Purpose", level=2)
    doc.add_paragraph("This document specifies the functional requirements for the Nexus E-Commerce "
                      "Platform API. It is intended for developers, testers and the change board.")
    doc.add_heading("1.2 Document Conventions", level=2)
    doc.add_paragraph("Requirement identifiers use the form REQ-NN. The word must indicates a mandatory "
                      "requirement.")
    doc.add_heading("1.3 Definitions, Acronyms and Abbreviations", level=2)
    _table(doc, ["Term", "Definition"], [
        ["JWT", "JSON Web Token, a signed bearer credential"],
        ["2FA", "Two-factor authentication"],
    ])
    doc.add_heading("1.4 References", level=2)
    doc.add_paragraph("IEEE Std 830-1998, Recommended Practice for Software Requirements Specifications.")

    doc.add_heading("2 Overall Description", level=1)
    doc.add_paragraph("The platform exposes a REST API consumed by the web storefront and mobile apps.")
    doc.add_picture(io.BytesIO(_png_bytes()))
    doc.add_paragraph("Figure 1: System context diagram", style="Caption")

    doc.add_heading("3 Specific Requirements", level=1)
    doc.add_heading("3.1 Core Functional Requirements", level=2)
    doc.add_paragraph("Table 2: Core functional requirements", style="Caption")
    priorities = ["High", "High", "Medium", "High", "High", "Medium"]
    table = _table(doc, ["ID", "Title", "Description", "Priority"],
                   [[r["id"], r["title"], r["body"], p] for r, p in zip(reqs[:6], priorities)])
    _mark_header_row(table)

    doc.add_heading("3.2 Order Lifecycle Requirements", level=2)
    for r in reqs[6:]:
        doc.add_heading(f"{r['id']}: {r['title']}", level=3)
        doc.add_paragraph(r["body"])

    doc.add_heading("Appendix A: Glossary", level=1)
    doc.add_paragraph("REQ-99: Example requirement format")
    doc.add_paragraph("Every requirement is written as an ID, a short title and a statement.")

    path = OUT / "ecommerce_srs_ieee.docx"
    doc.save(path)
    expected = [{"id": r["id"], "title": r["title"], "text": r["text"], "type": "functional",
                 "priority": priorities[i] if i < 6 else None} for i, r in enumerate(reqs)]
    _write_expected(path.name, expected, {"must_not_contain": [
        "Revision", "Contents", "Figure 1", "Table 2", "JSON Web Token", "IEEE Std", "REQ-99",
        "CONFIDENTIAL", "Page 1"]})
    return path


# ---------------------------------------------------------------------------
# Fixture 2 — banking SRS in key/value tables
# ---------------------------------------------------------------------------

def build_banking() -> Path:
    reqs = _md_requirements(BANKING_MD)
    doc = _new_document()
    _header_footer(doc, "Aurora Digital Banking - SRS", "Internal use only")
    _boilerplate_front(doc, "Aurora Digital Banking Platform")
    doc.add_heading("1 Introduction", level=1)
    doc.add_paragraph("Each requirement below is specified in its own requirement card.")
    doc.add_heading("2 Functional Requirements", level=1)
    for i, r in enumerate(reqs):
        doc.add_heading(f"2.{i + 1} {r['title']}", level=2)
        t = doc.add_table(rows=4, cols=2)
        t.style = "Table Grid"
        for row, (k, v) in zip(t.rows, [("Requirement ID", r["id"]), ("Title", r["title"]),
                                        ("Description", r["body"]), ("Priority", "High")]):
            row.cells[0].text = k
            row.cells[1].text = v
        doc.add_paragraph("")
    path = OUT / "banking_srs.docx"
    doc.save(path)
    _write_expected(path.name, [{"id": r["id"], "title": r["title"], "text": r["text"],
                                 "type": "functional", "priority": "High"} for r in reqs])
    return path


# ---------------------------------------------------------------------------
# Fixture 3 — messy real-world SRS
# ---------------------------------------------------------------------------

def build_messy() -> Path:
    doc = _new_document()
    _header_footer(doc, "Account Service SRS", "Page 2")
    doc.add_heading("Account Service - Software Requirements Specification", level=0)
    doc.add_heading("1 Introduction", level=1)
    doc.add_paragraph("This SRS was assembled from three team documents; formatting is inconsistent.")

    doc.add_heading("2 Functional Requirements", level=1)
    doc.add_paragraph("This section lists the functional requirements of the account service.")
    doc.add_paragraph("FR1: User Login")
    doc.add_paragraph("The system shall authenticate users with an email address and password.")
    doc.add_paragraph("FR-2 The system shall lock an account after five consecutive failed login attempts.")
    doc.add_paragraph("FR 3 - Password Reset")
    doc.add_paragraph("Users must be able to reset a forgotten password via a time-limited email link, as per FR-2.")

    doc.add_paragraph("[REQ-4] Session timeout")
    p = doc.add_paragraph()
    p._p.append(_run("Sessions shall expire after 30 minutes of inactivity"))
    p._p.append(_tracked("ins", " and require re-authentication", 1))
    p._p.append(_tracked("del", " silently", 2))
    p._p.append(_run("."))

    doc.add_heading("2.1 Account Data", level=2)
    header = ["Req ID", "Requirement", "Description", "Priority", "Acceptance Criteria"]
    t = doc.add_table(rows=8, cols=5)
    t.style = "Table Grid"
    for c, h in enumerate(header):
        t.rows[0].cells[c].text = h
    _bold_row(t)
    _mark_header_row(t)
    # row 1 — plain
    for c, v in enumerate(["FR-5.1", "Profile update",
                           "The system shall let users update their display name and avatar.",
                           "High", "Changes are visible after refresh."]):
        t.rows[1].cells[c].text = v
    # row 2 — multi-paragraph description cell
    for c, v in enumerate(["FR-5.2", "Data export", "", "Medium", ""]):
        t.rows[2].cells[c].text = v
    desc = t.rows[2].cells[2]
    desc.paragraphs[0].text = "The system shall export all personal data as JSON."
    desc.add_paragraph("The export must complete within 24 hours.")
    # row 3 — repeated header row (two tables pasted together)
    for c, h in enumerate(header):
        t.rows[3].cells[c].text = h
    # row 4 — section row merged across the whole table
    merged = t.rows[4].cells[0].merge(t.rows[4].cells[4])
    merged.text = "Deletion and audit"
    # row 5 — Requirement + Description merged horizontally
    t.rows[5].cells[0].text = "FR-5.3"
    rd = t.rows[5].cells[1].merge(t.rows[5].cells[2])
    rd.text = "The system shall permanently delete an account within 30 days of a verified request."
    t.rows[5].cells[3].text = "High"
    # rows 6-7 — ID, title and priority vertically merged; description continues
    for c in (0, 1, 3):
        t.rows[6].cells[c].merge(t.rows[7].cells[c])
    t.rows[6].cells[0].text = "FR-5.4"
    t.rows[6].cells[1].text = "Change audit trail"
    t.rows[6].cells[3].text = "Low"
    t.rows[6].cells[2].text = "The system shall keep an audit trail of profile changes."
    t.rows[7].cells[2].text = "Audit entries must be retained for one year."

    # 1x1 callout box holding a requirement
    box = doc.add_table(rows=1, cols=1)
    box.style = "Table Grid"
    cell = box.rows[0].cells[0]
    cell.paragraphs[0].text = "FR-6: Audit Logging"
    cell.add_paragraph("The system shall record every administrative action with a timestamp and actor.")

    # bold numbered "heading" typed as Normal text
    ph = doc.add_paragraph()
    r = ph.add_run("2.2 Notifications")
    r.bold = True

    # requirement inside a rich-text content control
    a = doc.add_paragraph("FR-7: Email Notifications")
    b = doc.add_paragraph("The system shall email users when their password changes.")
    _wrap_in_sdt([a, b], tag="requirement-block")

    # duplicate ID
    doc.add_paragraph("FR-2: Duplicate entry")
    doc.add_paragraph("This duplicate must be ignored by the parser.")

    doc.add_heading("3 Non-Functional Requirements", level=1)
    doc.add_heading("3.1 Performance", level=2)
    doc.add_paragraph("NFR-1: Response Time")
    doc.add_paragraph("Ninety-five percent of API requests shall complete within 300 milliseconds.")
    doc.add_heading("3.2 Security", level=2)
    doc.add_paragraph("NFR-2: Encryption")
    doc.add_paragraph("All personal data shall be encrypted at rest using AES-256.")

    doc.add_heading("Appendix A: Template", level=1)
    doc.add_paragraph("FR-99: Template example")
    doc.add_paragraph("This placeholder shows the requirement format and must not be extracted.")

    path = OUT / "messy_srs.docx"
    doc.save(path)
    expected = [
        {"id": "FR1", "title": "User Login", "type": "functional", "priority": None,
         "text": "User Login. The system shall authenticate users with an email address and password."},
        {"id": "FR-2", "title": "The system shall lock an account after five consecutive failed login attempts",
         "type": "functional", "priority": None,
         "text": "The system shall lock an account after five consecutive failed login attempts."},
        {"id": "FR-3", "title": "Password Reset", "type": "functional", "priority": None,
         "text": "Password Reset. Users must be able to reset a forgotten password via a time-limited "
                 "email link, as per FR-2."},
        {"id": "REQ-4", "title": "Session timeout", "type": "functional", "priority": None,
         "text": "Session timeout. Sessions shall expire after 30 minutes of inactivity and require "
                 "re-authentication."},
        {"id": "FR-5.1", "title": "Profile update", "type": "functional", "priority": "High",
         "text": "Profile update. The system shall let users update their display name and avatar. "
                 "Acceptance criteria: Changes are visible after refresh."},
        {"id": "FR-5.2", "title": "Data export", "type": "functional", "priority": "Medium",
         "text": "Data export. The system shall export all personal data as JSON. The export must "
                 "complete within 24 hours."},
        {"id": "FR-5.3", "title": "The system shall permanently delete an account within",
         "type": "functional", "priority": "High",
         "text": "The system shall permanently delete an account within 30 days of a verified request."},
        {"id": "FR-5.4", "title": "Change audit trail", "type": "functional", "priority": "Low",
         "text": "Change audit trail. The system shall keep an audit trail of profile changes. "
                 "Audit entries must be retained for one year."},
        {"id": "FR-6", "title": "Audit Logging", "type": "functional", "priority": None,
         "text": "Audit Logging. The system shall record every administrative action with a timestamp and actor."},
        {"id": "FR-7", "title": "Email Notifications", "type": "functional", "priority": None,
         "text": "Email Notifications. The system shall email users when their password changes."},
        {"id": "NFR-1", "title": "Response Time", "type": "non_functional", "priority": None,
         "text": "Response Time. Ninety-five percent of API requests shall complete within 300 milliseconds."},
        {"id": "NFR-2", "title": "Encryption", "type": "non_functional", "priority": None,
         "text": "Encryption. All personal data shall be encrypted at rest using AES-256."},
    ]
    _write_expected(path.name, expected, {
        "duplicates": ["FR-2"],
        "must_not_contain": ["silently", "Deletion and audit", "Req ID", "FR-99", "placeholder",
                             "Duplicate entry", "Account Service SRS"],
    })
    return path


# ---------------------------------------------------------------------------
# Fixture 4 — untagged SRS (shall-statement fallback)
# ---------------------------------------------------------------------------

def build_untagged() -> Path:
    doc = _new_document()
    doc.add_heading("Ordering Service SRS", level=0)
    doc.add_heading("1 Introduction", level=1)
    doc.add_paragraph("This document must be read by all stakeholders before implementation starts.")
    doc.add_heading("2 System Features", level=1)
    doc.add_heading("2.1 Ordering", level=2)
    doc.add_paragraph("Customers browse products and add them to a cart. The system shall allow customers to "
                      "place an order from the cart. The order confirmation must include an order number.")
    doc.add_heading("2.2 Payments", level=2)
    doc.add_paragraph("The system shall accept card payments.")
    path = OUT / "untagged_srs.docx"
    doc.save(path)
    sentences = [
        "The system shall allow customers to place an order from the cart.",
        "The order confirmation must include an order number.",
        "The system shall accept card payments.",
    ]
    _write_expected(path.name, [
        {"id": f"SRS-AUTO-{i:03d}", "title": s.rstrip("."), "text": s, "type": "functional", "priority": None}
        for i, s in enumerate(sentences, start=1)
    ], {"must_not_contain": ["stakeholders", "browse products"]})
    return path


# ---------------------------------------------------------------------------
# Fixture 5 — scanned (image-only) PDF
# ---------------------------------------------------------------------------

def build_scanned_pdf() -> Path:
    from PIL import Image, ImageDraw

    img = Image.new("RGB", (850, 1100), "white")
    d = ImageDraw.Draw(img)
    d.text((80, 80), "REQ-01: User Authentication", fill="black")
    d.text((80, 110), "The system must authenticate users.", fill="black")
    path = OUT / "scanned.pdf"
    img.save(path, format="PDF", resolution=100.0)
    return path


def export_pdf_with_word(docx_path: Path, pdf_path: Path) -> bool:
    """Word COM 'Save as PDF' (wdFormatPDF = 17). Windows + Microsoft Word only."""
    ps = (
        "$w = New-Object -ComObject Word.Application; $w.Visible = $false; "
        f"$d = $w.Documents.Open('{docx_path}', $false, $true); "
        f"$d.SaveAs([ref]'{pdf_path}', [ref]17); $d.Close($false); $w.Quit()"
    )
    try:
        subprocess.run(["powershell", "-NoProfile", "-Command", ps], check=True, timeout=180,
                       capture_output=True)
    except (OSError, subprocess.SubprocessError) as exc:
        print(f"PDF export failed: {exc}")
        return False
    return pdf_path.exists()


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--pdf", action="store_true", help="also export ecommerce_srs_ieee.pdf via Microsoft Word")
    args = ap.parse_args()
    OUT.mkdir(parents=True, exist_ok=True)
    built = [build_ecommerce(), build_banking(), build_messy(), build_untagged(), build_scanned_pdf()]
    if args.pdf:
        pdf = OUT / "ecommerce_srs_ieee.pdf"
        if export_pdf_with_word(built[0], pdf):
            built.append(pdf)
            expected = json.loads((OUT / "ecommerce_srs_ieee.expected.json").read_text(encoding="utf-8"))
            expected["fixture"] = pdf.name
            (OUT / "ecommerce_srs_ieee_pdf.expected.json").write_text(
                json.dumps(expected, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    for p in built:
        print(f"wrote {p.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
